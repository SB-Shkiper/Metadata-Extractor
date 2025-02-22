import sys
import os
from PyQt5.QtWidgets import QApplication, QWidget, QVBoxLayout, QPushButton, QLabel, QFileDialog, QTextEdit, QMessageBox
from PIL import Image
from PIL.ExifTags import TAGS
import PyPDF2
from docx import Document
from openpyxl import load_workbook
from moviepy.editor import VideoFileClip

class MetadataExtractor(QWidget):
    def __init__(self):
        super().__init__()
        self.initUI()

    def initUI(self):
        self.setWindowTitle('Metadata Extractor')
        self.setGeometry(100, 100, 600, 400)

        layout = QVBoxLayout()

        self.label = QLabel('Select a file to extract metadata:')
        layout.addWidget(self.label)

        self.button = QPushButton('Choose File')
        self.button.clicked.connect(self.openFileDialog)
        layout.addWidget(self.button)

        self.saveButton = QPushButton('Save as TXT')
        self.saveButton.clicked.connect(self.saveToFile)
        layout.addWidget(self.saveButton)

        self.textEdit = QTextEdit()
        self.textEdit.setReadOnly(True)  # Убираем возможность редактирования
        layout.addWidget(self.textEdit)

        self.setLayout(layout)

    def openFileDialog(self):
        options = QFileDialog.Options()
        fileName, _ = QFileDialog.getOpenFileName(self, "Select File", "", "Images (*.png *.jpg *.jpeg *.bmp *.tiff *.gif);;PDF Files (*.pdf);;Word Files (*.docx);;Excel Files (*.xlsx);;Video Files (*.mp4 *.avi *.mov);;All Files (*)", options=options)
        if fileName:
            self.extractMetadata(fileName)

    def extractMetadata(self, filePath):
        self.textEdit.clear()
        if filePath.lower().endswith(('.png', '.jpg', '.jpeg', '.bmp', '.tiff', '.gif')):
            self.extractImageMetadata(filePath)
        elif filePath.lower().endswith('.pdf'):
            self.extractPDFMetadata(filePath)
        elif filePath.lower().endswith('.docx'):
            self.extractWordMetadata(filePath)
        elif filePath.lower().endswith('.xlsx'):
            self.extractExcelMetadata(filePath)
        elif filePath.lower().endswith(('.mp4', '.avi', '.mov')):
            self.extractVideoMetadata(filePath)
        else:
            self.textEdit.setPlainText("Unsupported file format.")

    def extractImageMetadata(self, filePath):
        try:
            image = Image.open(filePath)
            exif_data = image._getexif()
            if exif_data:
                metadata = ""
                for tag_id, value in exif_data.items():
                    tag = TAGS.get(tag_id, tag_id)
                    metadata += f"{tag}: {value}\n"
                self.textEdit.setPlainText(metadata)
            else:
                self.textEdit.setPlainText("No EXIF data.")
        except Exception as e:
            self.textEdit.setPlainText(f"Error extracting metadata: {e}")

    def extractPDFMetadata(self, filePath):
        try:
            with open(filePath, 'rb') as file:
                reader = PyPDF2.PdfReader(file)
                metadata = reader.metadata
                if metadata:
                    metadata_str = "\n".join(f"{key}: {value}" for key, value in metadata.items())
                    self.textEdit.setPlainText(metadata_str)
                else:
                    self.textEdit.setPlainText("No metadata in PDF.")
        except Exception as e:
            self.textEdit.setPlainText(f"Error extracting metadata: {e}")

    def extractWordMetadata(self, filePath):
        try:
            doc = Document(filePath)
            metadata = doc.core_properties
            metadata_str = f"Title: {metadata.title}\nAuthor: {metadata.author}\nCreated: {metadata.created}"
            self.textEdit.setPlainText(metadata_str)
        except Exception as e:
            self.textEdit.setPlainText(f"Error extracting metadata: {e}")

    def extractExcelMetadata(self, filePath):
        try:
            wb = load_workbook(filePath)
            metadata = wb.properties
            metadata_str = f"Title: {metadata.title}\nAuthor: {metadata.creator}\nCreated: {metadata.created}"
            self.textEdit.setPlainText(metadata_str)
        except Exception as e:
            self.textEdit.setPlainText(f"Error extracting metadata: {e}")

    def extractVideoMetadata(self, filePath):
        try:
            clip = VideoFileClip(filePath)
            metadata_str = f"Duration: {clip.duration} seconds\n"
            metadata_str += f"FPS: {clip.fps}\n"
            metadata_str += f"Resolution: {clip .size[0]}x{clip.size[1]}"
            self.textEdit.setPlainText(metadata_str)
        except Exception as e:
            self.textEdit.setPlainText(f"Error extracting metadata: {e}")

    def saveToFile(self):
        options = QFileDialog.Options()
        fileName, _ = QFileDialog.getSaveFileName(self, "Save File", "", "Text Files (*.txt);;All Files (*)", options=options)
        if fileName:
            try:
                with open(fileName, 'w') as file:
                    file.write(self.textEdit.toPlainText())
                QMessageBox.information(self, "Success", "Metadata saved successfully.")
            except Exception as e:
                QMessageBox.warning(self, "Error", f"Could not save file: {e}")

if __name__ == '__main__':
    app = QApplication(sys.argv)
    extractor = MetadataExtractor()
    extractor.show()
    sys.exit(app.exec_())
